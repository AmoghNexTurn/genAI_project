import os
import uuid
import requests
import streamlit as st
from sentence_transformers import SentenceTransformer
from qdrant_client import QdrantClient
from qdrant_client.models import Distance, VectorParams, PointStruct
import pdfplumber
from docx import Document
import os
import json
import csv
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
import pandas as pd
from dotenv import load_dotenv, dotenv_values

print(1)

def sliding_window(text, window_size=1000, overlap=500):
    chunks = []
    start = 0
    while start < len(text):
        end = start + window_size
        chunks.append(text[start:end])
        if end >= len(text):
            break
        start += window_size - overlap
    return chunks

def chunk_file(file_path, row_chunk_size=10):
    ext = os.path.splitext(file_path)[1].lower()
    file_name = os.path.basename(file_path)
    chunks = []

    def make_payloads(text_chunks):
        return [{"text": chunk, "file": file_name} for chunk in text_chunks]

    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()
        chunks = make_payloads(sliding_window(text))

    elif ext == ".docx":
        doc = Document(file_path)
        full_text = "\n".join([p.text for p in doc.paragraphs])
        chunks = make_payloads(sliding_window(full_text))

    elif ext == ".pdf":
        with pdfplumber.open(file_path) as pdf:
            full_text = ""
            for page in pdf.pages:
                full_text += (page.extract_text() or "") + "\n"
        chunks = make_payloads(sliding_window(full_text))

    elif ext == ".json":
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list):
            text_chunks = [json.dumps(item, ensure_ascii=False) for item in data]
        elif isinstance(data, dict):
            text_chunks = [json.dumps({k: v}, ensure_ascii=False) for k, v in data.items()]
        else:
            text_chunks = [json.dumps(data, ensure_ascii=False)]
        chunks = make_payloads(text_chunks)

    elif ext == ".xml":
        tree = ET.parse(file_path)
        root = tree.getroot()
        text_chunks = [ET.tostring(elem, encoding='unicode') for elem in root]
        chunks = make_payloads(text_chunks)

    elif ext in [".csv", ".xlsx"]:
        if ext == ".csv":
            df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)

        text_chunks = []
        for i in range(0, len(df), row_chunk_size):
            chunk_df = df.iloc[i:i + row_chunk_size]
            text_chunks.append(chunk_df.to_csv(index=False, lineterminator='\n'))
        chunks = make_payloads(text_chunks)

    else:
        raise ValueError(f"Unsupported file extension: {ext}")

    return chunks

chunks = []
for file in os.listdir('../texts'):
    file_path = os.path.join('../texts', file)
    if os.path.isfile(file_path):
        try:
            chunks += chunk_file(file_path)
            print(f"File: {file}, Chunks: {len(chunks)}")
        except Exception as e:
            print(f"Error processing {file}: {e}")

model = SentenceTransformer('sentence-transformers/all-MiniLM-L6-v2')
embeddings = model.encode(chunks, convert_to_tensor=True)
print(f"Generated {len(embeddings)} embeddings.")

client = QdrantClient(host="localhost", port=6333)
collection_name = "battery_chunks"
vector_size = embeddings.shape[1]

client.recreate_collection(
    collection_name=collection_name,
    vectors_config=VectorParams(size=vector_size, distance=Distance.COSINE)
)

points = [
    PointStruct(id=i, vector=embeddings[i], payload=chunks[i])
    for i in range(len(chunks))
]

client.upsert(collection_name=collection_name, points=points)